#!/usr/bin/env python3
"""Validate DOCX and tagged-PDF exports for the 15 long-form guides."""

from __future__ import annotations

import argparse
import json
import math
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from docx import Document
from lxml import html
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[2]
SECTIONS_ROOT = ROOT / "content/themes/2026/sections"
DEFAULT_EXPORT_ROOT = ROOT / "app/content/regular-guides"
A4_WIDTH_POINTS = 595.276
A4_HEIGHT_POINTS = 841.89
A4_WIDTH_MM = 210.0
A4_HEIGHT_MM = 297.0
EMU_PER_MM = 36_000
RELATIONSHIP_NAMESPACE = {
    "r": "http://schemas.openxmlformats.org/package/2006/relationships"
}


def _load_sections() -> list[dict[str, object]]:
    sections: list[dict[str, object]] = []
    for metadata_path in SECTIONS_ROOT.glob("*/section.json"):
        metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
        if metadata.get("themeId") == "wsc-2026":
            sections.append(metadata)
    return sorted(sections, key=lambda section: int(section["order"]))


def _html_urls(slug: str) -> set[str]:
    source = (SECTIONS_ROOT / slug / "guide.html").read_text(encoding="utf-8")
    root = html.fragment_fromstring(source, create_parent="div")
    return {anchor.get("href") for anchor in root.xpath(".//a[@href]")}


def _docx_urls(path: Path) -> set[str]:
    with zipfile.ZipFile(path) as archive:
        relationships = ElementTree.fromstring(
            archive.read("word/_rels/document.xml.rels")
        )
    return {
        relationship.get("Target")
        for relationship in relationships.findall("r:Relationship", RELATIONSHIP_NAMESPACE)
        if relationship.get("TargetMode") == "External"
    }


def _pdf_urls(reader: PdfReader) -> tuple[set[str], list[str]]:
    urls: set[str] = set()
    errors: list[str] = []
    for page_index, page in enumerate(reader.pages, start=1):
        media_box = page.mediabox
        left, bottom = float(media_box.left), float(media_box.bottom)
        right, top = float(media_box.right), float(media_box.top)
        for annotation_reference in page.get("/Annots", []):
            annotation = annotation_reference.get_object()
            if annotation.get("/Subtype") != "/Link":
                continue
            action = annotation.get("/A")
            if action and action.get("/URI"):
                urls.add(str(action.get("/URI")))
            rectangle = annotation.get("/Rect")
            if rectangle and len(rectangle) == 4:
                x1, y1, x2, y2 = map(float, rectangle)
                if (
                    min(x1, x2) < left - 0.5
                    or max(x1, x2) > right + 0.5
                    or min(y1, y2) < bottom - 0.5
                    or max(y1, y2) > top + 0.5
                ):
                    errors.append(
                        f"page {page_index} contains an out-of-bounds link rectangle"
                    )
    return urls, errors


def _outline_count(value: object) -> int:
    if isinstance(value, list):
        return sum(_outline_count(item) for item in value)
    return 1


def _close(left: float, right: float, tolerance: float = 0.75) -> bool:
    return math.isclose(left, right, abs_tol=tolerance)


def _validate_export(
    section: dict[str, object], export_root: Path
) -> tuple[list[str], dict[str, object]]:
    slug = str(section["id"])
    errors: list[str] = []
    docx_path = export_root / "docx" / f"{slug}.docx"
    pdf_path = export_root / "pdf" / f"{slug}.pdf"
    for path in (docx_path, pdf_path):
        if not path.exists():
            return [f"missing export: {path}"], {"slug": slug}

    html_urls = _html_urls(slug)
    docx_urls = _docx_urls(docx_path)
    if docx_urls != html_urls:
        errors.append(
            "DOCX URL set differs from HTML: "
            f"missing={sorted(html_urls - docx_urls)}, extra={sorted(docx_urls - html_urls)}"
        )

    document = Document(docx_path)
    if document.tables:
        errors.append(f"DOCX contains {len(document.tables)} table(s)")
    for section_index, docx_section in enumerate(document.sections, start=1):
        width_mm = float(docx_section.page_width) / EMU_PER_MM
        height_mm = float(docx_section.page_height) / EMU_PER_MM
        if not (_close(width_mm, A4_WIDTH_MM, 0.1) and _close(height_mm, A4_HEIGHT_MM, 0.1)):
            errors.append(
                f"DOCX section {section_index} is {width_mm:.2f} x {height_mm:.2f} mm, not A4"
            )

    reader = PdfReader(pdf_path)
    pdf_urls, pdf_link_errors = _pdf_urls(reader)
    errors.extend(pdf_link_errors)
    if pdf_urls != html_urls:
        errors.append(
            "PDF URL set differs from HTML: "
            f"missing={sorted(html_urls - pdf_urls)}, extra={sorted(pdf_urls - html_urls)}"
        )

    root = reader.trailer["/Root"]
    tagged = root.get("/StructTreeRoot") is not None
    if not tagged:
        errors.append("PDF has no structure tree")

    outlines = reader.outline
    outline_count = _outline_count(outlines) if outlines else 0
    if outline_count == 0:
        errors.append("PDF has no bookmarks")

    for page_index, page in enumerate(reader.pages, start=1):
        width = float(page.mediabox.width)
        height = float(page.mediabox.height)
        if not (
            _close(width, A4_WIDTH_POINTS)
            and _close(height, A4_HEIGHT_POINTS)
        ):
            errors.append(
                f"PDF page {page_index} is {width:.2f} x {height:.2f} points, not A4"
            )
        if not (page.extract_text() or "").strip():
            errors.append(f"PDF page {page_index} has no extractable text")

    expected_title = str(section["title"])
    metadata_title = (reader.metadata.title or "").strip() if reader.metadata else ""
    normalized_expected = expected_title.replace("’", "'")
    normalized_actual = metadata_title.replace("’", "'")
    if normalized_actual != normalized_expected:
        errors.append(
            f"PDF title metadata is {metadata_title!r}; expected {expected_title!r}"
        )

    metrics = {
        "slug": slug,
        "pages": len(reader.pages),
        "links": len(html_urls),
        "bookmarks": outline_count,
        "tagged": tagged,
        "docxBytes": docx_path.stat().st_size,
        "pdfBytes": pdf_path.stat().st_size,
    }
    return errors, metrics


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--export-root",
        type=Path,
        default=DEFAULT_EXPORT_ROOT,
        help="Root containing docx/ and pdf/ guide exports.",
    )
    parser.add_argument(
        "--section",
        action="append",
        help="Validate only this slug; repeat to select several sections.",
    )
    parser.add_argument("--json", action="store_true", help="Print JSON metrics.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    sections = _load_sections()
    if args.section:
        wanted = set(args.section)
        sections = [section for section in sections if section["id"] in wanted]
        unknown = wanted - {str(section["id"]) for section in sections}
        if unknown:
            print("Unknown section slug(s): " + ", ".join(sorted(unknown)), file=sys.stderr)
            return 2

    all_errors: list[tuple[str, str]] = []
    all_metrics: list[dict[str, object]] = []
    for section in sections:
        errors, metrics = _validate_export(section, args.export_root.resolve())
        all_metrics.append(metrics)
        for error in errors:
            all_errors.append((str(section["id"]), error))

    if args.json:
        print(json.dumps({"guides": all_metrics, "errors": all_errors}, indent=2))
    else:
        for metrics in all_metrics:
            print(
                f"{metrics['slug']}: {metrics.get('pages', 0)} pages, "
                f"{metrics.get('links', 0)} unique links, "
                f"{metrics.get('bookmarks', 0)} bookmarks, "
                f"tagged={metrics.get('tagged', False)}"
            )
        for slug, error in all_errors:
            print(f"ERROR {slug}: {error}", file=sys.stderr)

    if all_errors:
        print(f"Export validation failed with {len(all_errors)} error(s).", file=sys.stderr)
        return 1
    if not args.json:
        print(f"Validated {len(all_metrics)} DOCX/PDF guide pair(s).")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
