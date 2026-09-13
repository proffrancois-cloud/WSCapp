#!/usr/bin/env python3
"""Build the long-form WSCapp guides from their canonical HTML sources.

The generated DOCX files are editable delivery copies, not source files.
LibreOffice exports them to PDF so headings, lists, links, bookmarks, and
reading order survive as far as the DOCX/PDF toolchain permits.
"""

from __future__ import annotations

import argparse
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterator

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Mm, Pt, RGBColor
from lxml import html
from lxml.html import HtmlElement


ROOT = Path(__file__).resolve().parents[2]
THEME_SECTIONS = ROOT / "content/themes/2026/sections"
DEFAULT_OUTPUT_ROOT = ROOT / "app/content/regular-guides"
GUIDE_MASCOT = ROOT / "app/assets/mascot/library/final-pack/Guide.png"
BUNDLED_SOFFICE = (
    Path.home()
    / ".cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/soffice"
)

# WSCapp's warm reading palette. These values intentionally match the approved
# PDF reskin rather than Word's blue default theme.
CREAM = "FFFAF2"
PALE_GOLD = "F7EACC"
QUESTION_GOLD = "F5E5B8"
PALE_ROSE = "F4E5D9"
COCOA = "412A17"
OLIVE = "6B4E16"
GOLD = "C89C44"
DARK_GOLD = "825F1F"
MUTED_BROWN = "6D5238"
BODY_FONT = "Georgia"
DISPLAY_FONT = "Avenir Next"
FIXED_TIMESTAMP = datetime(2026, 7, 14, 0, 0, 0)
ZIP_TIMESTAMP = (2026, 7, 14, 0, 0, 0)


@dataclass(frozen=True)
class GuideConfig:
    slug: str
    subtitle: str
    footer_title: str


GUIDES = {
    "introductory-questions": GuideConfig(
        slug="introductory-questions",
        subtitle=(
            "How journeys, destinations, thresholds, and endings shape the "
            "question at the heart of the 2026 curriculum"
        ),
        footer_title="INTRODUCTORY QUESTIONS",
    ),
    "progress-not-regress": GuideConfig(
        slug="progress-not-regress",
        subtitle=(
            "How people measure progress, endure uncertainty, and decide "
            "whether movement is real"
        ),
        footer_title="PROGRESS, NOT REGRESS",
    ),
    "more-to-do-than-can-ever-be-listed": GuideConfig(
        slug="more-to-do-than-can-ever-be-listed",
        subtitle=(
            "Lists, unfinished tasks, priorities, and the systems people use "
            "to make work feel possible"
        ),
        footer_title="MORE TO DO THAN CAN BE LISTED",
    ),
    "the-end-is-nearish": GuideConfig(
        slug="the-end-is-nearish",
        subtitle=(
            "Caretakers, successors, final acts, and the unstable interval "
            "between ending and replacement"
        ),
        footer_title="THE END IS NEARISH",
    ),
    "theres-a-draft-in-here": GuideConfig(
        slug="theres-a-draft-in-here",
        subtitle=(
            "How sketches, demos, cuts, and rehearsal turn unfinished "
            "work into better decisions"
        ),
        footer_title="THERE'S A DRAFT IN HERE",
    ),
    "were-all-in-this-to-get-there": GuideConfig(
        slug="were-all-in-this-to-get-there",
        subtitle=(
            "Why shared destinations demand coordination, leadership, "
            "planning, and honest measures of progress"
        ),
        footer_title="WE'RE ALL IN THIS TO GET THERE",
    ),
    "where-the-sidewalk-starts": GuideConfig(
        slug="where-the-sidewalk-starts",
        subtitle=(
            "Sidewalks, parking lots, public space, and the overlooked "
            "places between departure and arrival"
        ),
        footer_title="WHERE THE SIDEWALK STARTS",
    ),
    "monkey-see-monkey-prototype": GuideConfig(
        slug="monkey-see-monkey-prototype",
        subtitle=(
            "Prototypes, experiments, users, and the difficult path from "
            "promising model to working reality"
        ),
        footer_title="MONKEY SEE, MONKEY PROTOTYPE",
    ),
    "the-lovely-and-the-liminal": GuideConfig(
        slug="the-lovely-and-the-liminal",
        subtitle=(
            "Doorways, backrooms, rituals, and the strange power of being "
            "neither here nor there"
        ),
        footer_title="THE LOVELY AND THE LIMINAL",
    ),
    "going-pains": GuideConfig(
        slug="going-pains",
        subtitle=(
            "Adolescence, separation, migration, and the costs hidden inside "
            "growth and transition"
        ),
        footer_title="GOING PAINS",
    ),
    "home-and-wandering": GuideConfig(
        slug="home-and-wandering",
        subtitle=(
            "Navigation, migration, wandering, and the networks that let "
            "travellers find both routes and refuge"
        ),
        footer_title="HOME AND WANDERING",
    ),
    "where-were-going-well-still-need-them": GuideConfig(
        slug="where-were-going-well-still-need-them",
        subtitle=(
            "Roads, routes, infrastructure, and why every future still "
            "depends on choices made on the ground"
        ),
        footer_title="WHERE WE'RE GOING",
    ),
    "call-of-duty-free": GuideConfig(
        slug="call-of-duty-free",
        subtitle=(
            "Travel, tourism, hospitality, and the networks that move "
            "people and goods"
        ),
        footer_title="CALL OF DUTY-FREE",
    ),
    "next-year-in-futurism": GuideConfig(
        slug="next-year-in-futurism",
        subtitle=(
            "Predictions, perpetual promises, and the hard work of telling "
            "a plausible future from an attractive fantasy"
        ),
        footer_title="NEXT YEAR IN FUTURISM",
    ),
    "concluding-questions": GuideConfig(
        slug="concluding-questions",
        subtitle=(
            "Patience, surrender, infinity, and the final questions that test "
            "what it means to arrive"
        ),
        footer_title="CONCLUDING QUESTIONS",
    ),
}

GUIDE_ALIASES = {
    "call": "call-of-duty-free",
    "draft": "theres-a-draft-in-here",
}


@dataclass(frozen=True)
class InlineSegment:
    text: str
    bold: bool = False
    italic: bool = False
    url: str | None = None


def _set_paragraph_border(
    paragraph: object,
    *,
    side: str,
    color: str,
    size: int = 12,
    space: int = 4,
) -> None:
    properties = paragraph._p.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        properties.append(borders)
    border = borders.find(qn(f"w:{side}"))
    if border is None:
        border = OxmlElement(f"w:{side}")
        borders.append(border)
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), str(size))
    border.set(qn("w:space"), str(space))
    border.set(qn("w:color"), color)


def _set_paragraph_shading(paragraph: object, fill: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")


def _set_style_font(style: object, name: str, size: float, color: str) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    run_properties = style.element.get_or_add_rPr()
    fonts = run_properties.get_or_add_rFonts()
    for script in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{script}"), name)
    language = run_properties.find(qn("w:lang"))
    if language is None:
        language = OxmlElement("w:lang")
        run_properties.append(language)
    language.set(qn("w:val"), "en-US")


def _set_outline_level(style: object, level: int) -> None:
    properties = style.element.get_or_add_pPr()
    outline = properties.find(qn("w:outlineLvl"))
    if outline is None:
        outline = OxmlElement("w:outlineLvl")
        properties.append(outline)
    outline.set(qn("w:val"), str(level))


def _configure_styles(document: DocumentType) -> None:
    styles = document.styles

    normal = styles["Normal"]
    _set_style_font(normal, BODY_FONT, 10.2, COCOA)
    normal.paragraph_format.line_spacing = 1.19
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.widow_control = True

    title = styles["Title"]
    _set_style_font(title, DISPLAY_FONT, 30, COCOA)
    title.font.bold = True
    title.paragraph_format.space_before = Pt(4)
    title.paragraph_format.space_after = Pt(5)
    title.paragraph_format.keep_with_next = True
    # The stock Word template draws an Accent 1 (blue) rule below Title.
    # Remove it so the only rule on the opening page is WSCapp gold.
    title_properties = title.element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)

    heading_specs = {
        "Heading 1": (DISPLAY_FONT, 19, OLIVE, 14, 9, 0),
        "Heading 2": (DISPLAY_FONT, 12.4, COCOA, 12, 4, 1),
        "Heading 3": (DISPLAY_FONT, 10.8, OLIVE, 9, 3, 2),
    }
    for style_name, (font, size, color, before, after, outline) in heading_specs.items():
        style = styles[style_name]
        _set_style_font(style, font, size, color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        _set_outline_level(style, outline)

    # Make lists read like paragraphs in a book, with enough hanging indent to
    # scan but without the cramped default Word appearance.
    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        _set_style_font(style, BODY_FONT, 10.2, COCOA)
        style.paragraph_format.left_indent = Mm(7)
        style.paragraph_format.first_line_indent = Mm(-3.4)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.2

    block_quote = styles["Quote"]
    _set_style_font(block_quote, BODY_FONT, 11.0, OLIVE)
    block_quote.font.italic = True
    block_quote.paragraph_format.left_indent = Mm(7)
    block_quote.paragraph_format.right_indent = Mm(5)
    block_quote.paragraph_format.space_before = Pt(8)
    block_quote.paragraph_format.space_after = Pt(9)
    block_quote.paragraph_format.line_spacing = 1.2

    footer = styles["Footer"]
    _set_style_font(footer, DISPLAY_FONT, 7.2, MUTED_BROWN)

    def add_paragraph_style(name: str, base: str = "Normal") -> object:
        try:
            return styles[name]
        except KeyError:
            style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = styles[base]
            return style

    kicker = add_paragraph_style("WSC Kicker")
    _set_style_font(kicker, DISPLAY_FONT, 8.5, DARK_GOLD)
    kicker.font.bold = True
    kicker.paragraph_format.space_after = Pt(7)
    kicker.paragraph_format.keep_with_next = True

    subtitle = add_paragraph_style("WSC Subtitle")
    _set_style_font(subtitle, BODY_FONT, 12.0, MUTED_BROWN)
    subtitle.font.italic = True
    subtitle.paragraph_format.space_after = Pt(18)
    subtitle.paragraph_format.keep_with_next = True

    entry_lead = add_paragraph_style("WSC Entry Lead")
    _set_style_font(entry_lead, DISPLAY_FONT, 10.7, COCOA)
    entry_lead.font.bold = True
    entry_lead.paragraph_format.space_before = Pt(9)
    entry_lead.paragraph_format.space_after = Pt(3)
    entry_lead.paragraph_format.keep_with_next = True

    entry_detail = add_paragraph_style("WSC Entry Detail")
    _set_style_font(entry_detail, BODY_FONT, 10.2, COCOA)
    entry_detail.paragraph_format.left_indent = Mm(3.5)
    entry_detail.paragraph_format.space_after = Pt(4)
    entry_detail.paragraph_format.line_spacing = 1.19

    overview = add_paragraph_style("WSC Overview Entry")
    _set_style_font(overview, BODY_FONT, 9.25, COCOA)
    overview.paragraph_format.left_indent = Mm(2)
    overview.paragraph_format.first_line_indent = Mm(-2)
    overview.paragraph_format.space_after = Pt(3.5)
    overview.paragraph_format.line_spacing = 1.08
    overview.paragraph_format.keep_together = True

    source_note = add_paragraph_style("WSC Source Note")
    _set_style_font(source_note, BODY_FONT, 8.5, MUTED_BROWN)
    source_note.paragraph_format.left_indent = Mm(3)
    source_note.paragraph_format.right_indent = Mm(1)
    source_note.paragraph_format.space_before = Pt(4)
    source_note.paragraph_format.space_after = Pt(5)
    source_note.paragraph_format.line_spacing = 1.08

    pitfall = add_paragraph_style("WSC Pitfall")
    _set_style_font(pitfall, BODY_FONT, 9.5, COCOA)
    pitfall.paragraph_format.left_indent = Mm(4)
    pitfall.paragraph_format.right_indent = Mm(2)
    pitfall.paragraph_format.space_before = Pt(7)
    pitfall.paragraph_format.space_after = Pt(8)
    pitfall.paragraph_format.line_spacing = 1.18

    question = add_paragraph_style("WSC Study Question")
    _set_style_font(question, BODY_FONT, 10.0, COCOA)
    question.font.italic = True
    question.paragraph_format.left_indent = Mm(4)
    question.paragraph_format.right_indent = Mm(2)
    question.paragraph_format.space_before = Pt(8)
    question.paragraph_format.space_after = Pt(10)
    question.paragraph_format.line_spacing = 1.18

    try:
        link = styles["WSC Link"]
    except KeyError:
        link = styles.add_style("WSC Link", WD_STYLE_TYPE.CHARACTER)
    link.font.name = BODY_FONT
    link.font.color.rgb = RGBColor.from_string(DARK_GOLD)
    link.font.underline = True
    link_fonts = link.element.get_or_add_rPr().get_or_add_rFonts()
    for script in ("ascii", "hAnsi", "eastAsia", "cs"):
        link_fonts.set(qn(f"w:{script}"), BODY_FONT)


def _add_page_background(document: DocumentType) -> None:
    background = document._element.find(qn("w:background"))
    if background is None:
        background = OxmlElement("w:background")
        document._element.insert(0, background)
    background.set(qn("w:color"), CREAM)

    settings = document.settings.element
    if settings.find(qn("w:displayBackgroundShape")) is None:
        settings.append(OxmlElement("w:displayBackgroundShape"))
    if settings.find(qn("w:updateFields")) is None:
        update_fields = OxmlElement("w:updateFields")
        update_fields.set(qn("w:val"), "true")
        settings.append(update_fields)


def _add_field(paragraph: object, instruction: str) -> None:
    begin_run = OxmlElement("w:r")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run.append(begin)

    instruction_run = OxmlElement("w:r")
    instruction_text = OxmlElement("w:instrText")
    instruction_text.set(qn("xml:space"), "preserve")
    instruction_text.text = f" {instruction} "
    instruction_run.append(instruction_text)

    separate_run = OxmlElement("w:r")
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    separate_run.append(separate)

    result_run = OxmlElement("w:r")
    result_text = OxmlElement("w:t")
    result_text.text = "1"
    result_run.append(result_text)

    end_run = OxmlElement("w:r")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run.append(end)

    for element in (begin_run, instruction_run, separate_run, result_run, end_run):
        paragraph._p.append(element)


def _format_footer(footer: object, section: object, title: str) -> None:
    paragraph = footer.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        section.page_width - section.left_margin - section.right_margin,
        WD_TAB_ALIGNMENT.RIGHT,
    )
    _set_paragraph_border(paragraph, side="top", color=GOLD, size=8, space=4)

    left = paragraph.add_run(title)
    left.font.name = DISPLAY_FONT
    left.font.size = Pt(7.2)
    left.font.bold = True
    left.font.color.rgb = RGBColor.from_string(MUTED_BROWN)
    left._element.get_or_add_rPr().get_or_add_rFonts().set(
        qn("w:eastAsia"), DISPLAY_FONT
    )

    page_label = paragraph.add_run("\tPAGE ")
    page_label.font.name = DISPLAY_FONT
    page_label.font.size = Pt(7.2)
    page_label.font.color.rgb = RGBColor.from_string(MUTED_BROWN)
    _add_field(paragraph, "PAGE")
    of_label = paragraph.add_run(" OF ")
    of_label.font.name = DISPLAY_FONT
    of_label.font.size = Pt(7.2)
    of_label.font.color.rgb = RGBColor.from_string(MUTED_BROWN)
    _add_field(paragraph, "NUMPAGES")


def _configure_page(document: DocumentType, config: GuideConfig) -> None:
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(20)
    section.right_margin = Mm(20)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(19)
    section.header_distance = Mm(6)
    section.footer_distance = Mm(8)
    section.different_first_page_header_footer = True

    first_header = section.first_page_header
    first_header.is_linked_to_previous = False
    header_paragraph = first_header.paragraphs[0]
    header_paragraph.clear()
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_paragraph.paragraph_format.space_after = Pt(0)
    image = header_paragraph.add_run().add_picture(
        str(GUIDE_MASCOT), width=Inches(0.64)
    )
    image._inline.docPr.set("title", "WSCapp Guide mascot")
    image._inline.docPr.set("descr", "Guide, the WSCapp guide mascot")

    _format_footer(section.footer, section, config.footer_title)
    _format_footer(section.first_page_footer, section, config.footer_title)


def _paragraph_class(element: HtmlElement, forced_style: str | None = None) -> str:
    if forced_style:
        return forced_style
    classes = set((element.get("class") or "").split())
    class_text = " ".join(classes).lower()
    text = " ".join(element.text_content().split()).lower()
    if "book-deck" in classes:
        return "WSC Subtitle"
    if "source-note" in class_text or text.startswith(
        ("source:", "sources:", "required source trail:", "evidence trail:")
    ):
        return "WSC Source Note"
    if "pitfall" in class_text or text.startswith("pitfall:"):
        return "WSC Pitfall"
    if (
        "mini-question" in class_text
        or "study-question" in class_text
        or text.startswith("mini-question:")
    ):
        return "WSC Study Question"
    return "Normal"


def _iter_inline_segments(
    node: HtmlElement,
    *,
    bold: bool = False,
    italic: bool = False,
    url: str | None = None,
) -> Iterator[InlineSegment]:
    if node.text:
        yield InlineSegment(node.text, bold=bold, italic=italic, url=url)

    for child in node:
        tag = child.tag.lower() if isinstance(child.tag, str) else ""
        child_bold = bold or tag in {"strong", "b"}
        child_italic = italic or tag in {"em", "i", "cite"}
        child_url = child.get("href") if tag == "a" else url
        if tag == "br":
            yield InlineSegment("\n", bold=bold, italic=italic, url=url)
        else:
            yield from _iter_inline_segments(
                child,
                bold=child_bold,
                italic=child_italic,
                url=child_url,
            )
        if child.tail:
            yield InlineSegment(child.tail, bold=bold, italic=italic, url=url)


def _normalized_segments(node: HtmlElement) -> list[InlineSegment]:
    segments: list[InlineSegment] = []
    for segment in _iter_inline_segments(node):
        text = re.sub(r"[\t\r\f\v ]+", " ", segment.text)
        text = re.sub(r" *\n *", "\n", text)
        if not text:
            continue
        if not segments:
            text = text.lstrip()
        elif segments[-1].text.endswith((" ", "\n")) and text.startswith(" "):
            text = text.lstrip(" ")
        if text:
            segments.append(
                InlineSegment(
                    text=text,
                    bold=segment.bold,
                    italic=segment.italic,
                    url=segment.url,
                )
            )
    if segments:
        last = segments[-1]
        segments[-1] = InlineSegment(
            text=last.text.rstrip(),
            bold=last.bold,
            italic=last.italic,
            url=last.url,
        )
    return [segment for segment in segments if segment.text]


def _add_hyperlink(
    paragraph: object,
    text: str,
    url: str,
    *,
    bold: bool = False,
    italic: bool = False,
) -> None:
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)

    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    run_style = OxmlElement("w:rStyle")
    run_style.set(qn("w:val"), "WSCLink")
    run_properties.append(run_style)

    fonts = OxmlElement("w:rFonts")
    for script in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{script}"), BODY_FONT)
    run_properties.append(fonts)

    color = OxmlElement("w:color")
    color.set(qn("w:val"), DARK_GOLD)
    run_properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(underline)
    if bold:
        run_properties.append(OxmlElement("w:b"))
    if italic:
        run_properties.append(OxmlElement("w:i"))

    text_element = OxmlElement("w:t")
    if text.startswith(" ") or text.endswith(" "):
        text_element.set(qn("xml:space"), "preserve")
    text_element.text = text
    run.extend((run_properties, text_element))
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _append_inline(paragraph: object, node: HtmlElement) -> None:
    for segment in _normalized_segments(node):
        if segment.url:
            _add_hyperlink(
                paragraph,
                segment.text,
                segment.url,
                bold=segment.bold,
                italic=segment.italic,
            )
            continue
        run = paragraph.add_run(segment.text)
        run.bold = segment.bold
        run.italic = segment.italic


def _decorate_callout(paragraph: object, style_name: str) -> None:
    if style_name == "WSC Subtitle":
        _set_paragraph_border(paragraph, side="bottom", color=GOLD, size=10, space=8)
    elif style_name == "WSC Source Note":
        _set_paragraph_shading(paragraph, PALE_GOLD)
        _set_paragraph_border(paragraph, side="left", color=GOLD, size=14, space=6)
    elif style_name == "WSC Pitfall":
        _set_paragraph_shading(paragraph, PALE_ROSE)
        _set_paragraph_border(paragraph, side="left", color=OLIVE, size=14, space=6)
    elif style_name == "WSC Study Question":
        _set_paragraph_shading(paragraph, QUESTION_GOLD)
        _set_paragraph_border(paragraph, side="left", color=DARK_GOLD, size=14, space=6)


def _add_html_paragraph(
    document: DocumentType,
    element: HtmlElement,
    *,
    forced_style: str | None = None,
) -> object:
    style_name = _paragraph_class(element, forced_style)
    paragraph = document.add_paragraph(style=style_name)
    _append_inline(paragraph, element)
    _decorate_callout(paragraph, style_name)
    return paragraph


def _add_list(document: DocumentType, element: HtmlElement, *, level: int = 0) -> None:
    style_name = "List Number" if element.tag.lower() == "ol" else "List Bullet"
    for item in element.xpath("./li"):
        paragraph = document.add_paragraph(style=style_name)
        if level:
            paragraph.paragraph_format.left_indent = Mm(7 + level * 5)
        # Only the inline content that belongs directly to this list item is
        # added here; nested lists become their own paragraphs below it.
        shallow = html.Element("span")
        shallow.text = item.text
        for child in item:
            if isinstance(child.tag, str) and child.tag.lower() in {"ul", "ol"}:
                continue
            shallow.append(html.fromstring(html.tostring(child, encoding="unicode")))
        _append_inline(paragraph, shallow)
        for nested in item.xpath("./ul|./ol"):
            _add_list(document, nested, level=level + 1)


def _cell_text(cell: HtmlElement) -> str:
    return " ".join(cell.text_content().split())


def _add_table_as_prose(
    document: DocumentType,
    table: HtmlElement,
    *,
    compact_overview: bool = False,
) -> None:
    """Turn a source table into short subsections instead of a DOCX table.

    The source HTML uses tables as a compact drafting device. In the reading
    edition, the first column becomes an entry lead and each remaining column
    becomes a labelled paragraph. This retains every fact while producing the
    continuous, book-like page the app needs.
    """

    rows = table.xpath(".//tr")
    if not rows:
        return

    header_cells = rows[0].xpath("./th|./td")
    headers = [_cell_text(cell) for cell in header_cells]
    data_rows = rows[1:] if rows[0].xpath("./th") else rows

    for row_index, row in enumerate(data_rows):
        cells = row.xpath("./th|./td")
        if not cells:
            continue
        lead_text = _cell_text(cells[0])
        if not lead_text:
            continue

        if compact_overview:
            overview = document.add_paragraph(style="WSC Overview Entry")
            lead = overview.add_run(lead_text.rstrip(".:") + ". ")
            lead.bold = True
            lead.font.color.rgb = RGBColor.from_string(OLIVE)
            for column_index, cell in enumerate(cells[1:], start=1):
                if not _cell_text(cell):
                    continue
                if column_index < len(headers) and headers[column_index]:
                    label = overview.add_run(headers[column_index].rstrip(".:") + ": ")
                    label.bold = True
                    label.font.color.rgb = RGBColor.from_string(MUTED_BROWN)
                _append_inline(overview, cell)
                if column_index < len(cells) - 1:
                    overview.add_run(" ")
            continue

        if len(lead_text) <= 92 and len(lead_text.split()) <= 13:
            lead = document.add_paragraph(style="Heading 2")
        else:
            lead = document.add_paragraph(style="WSC Entry Lead")
        _append_inline(lead, cells[0])

        for column_index, cell in enumerate(cells[1:], start=1):
            if not _cell_text(cell):
                continue
            detail = document.add_paragraph(style="WSC Entry Detail")
            if column_index < len(headers) and headers[column_index]:
                label = detail.add_run(headers[column_index].rstrip(".:") + ". ")
                label.bold = True
                label.font.color.rgb = RGBColor.from_string(OLIVE)
            _append_inline(detail, cell)

        # A little breathing room between entries without drawing boxes or
        # rules that would make the section look like a worksheet.
        if row_index < len(data_rows) - 1 and document.paragraphs:
            document.paragraphs[-1].paragraph_format.space_after = Pt(7)


def _add_blockquote(document: DocumentType, element: HtmlElement) -> None:
    children = element.xpath("./p")
    if children:
        for child in children:
            _add_html_paragraph(document, child, forced_style="Quote")
        return
    paragraph = document.add_paragraph(style="Quote")
    _append_inline(paragraph, element)


def _walk_block(
    document: DocumentType,
    element: HtmlElement,
    *,
    h3_count: list[int],
    forced_style: str | None = None,
) -> None:
    tag = element.tag.lower() if isinstance(element.tag, str) else ""
    classes = set((element.get("class") or "").split())

    if tag == "h2":
        document.add_paragraph("WSCAPP 2026 | SCHOLAR'S GUIDE", style="WSC Kicker")
        title = document.add_paragraph(style="Title")
        _append_inline(title, element)
        return

    if tag == "h3":
        paragraph = document.add_paragraph(style="Heading 1")
        if h3_count[0] > 0 and "book-conclusion" not in classes:
            paragraph.paragraph_format.page_break_before = True
        h3_count[0] += 1
        _append_inline(paragraph, element)
        _set_paragraph_border(paragraph, side="bottom", color=GOLD, size=9, space=5)
        return

    if tag in {"h4", "h5"}:
        style_name = "Heading 2" if tag == "h4" else "Heading 3"
        paragraph = document.add_paragraph(style=style_name)
        _append_inline(paragraph, element)
        return

    if tag == "p":
        _add_html_paragraph(document, element, forced_style=forced_style)
        return

    if tag in {"ul", "ol"}:
        _add_list(document, element)
        return

    if tag == "blockquote":
        _add_blockquote(document, element)
        return

    if tag == "table":
        _add_table_as_prose(
            document,
            element,
            compact_overview=h3_count[0] == 1,
        )
        return

    if tag in {"article", "main", "section", "div"}:
        table = element.xpath("./table")
        if table:
            for child_table in table:
                _add_table_as_prose(
                    document,
                    child_table,
                    compact_overview=h3_count[0] == 1,
                )
            return
        next_style = forced_style
        class_text = " ".join(classes).lower()
        if "source-note" in class_text:
            next_style = "WSC Source Note"
        elif "pitfall" in class_text:
            next_style = "WSC Pitfall"
        elif "mini-question" in class_text or "study-question" in class_text:
            next_style = "WSC Study Question"
        for child in element:
            _walk_block(
                document,
                child,
                h3_count=h3_count,
                forced_style=next_style,
            )


def _remove_initial_empty_paragraph(document: DocumentType) -> None:
    if len(document.paragraphs) != 1 or document.paragraphs[0].text:
        return
    paragraph = document.paragraphs[0]._element
    paragraph.getparent().remove(paragraph)


def _add_subtitle(document: DocumentType, config: GuideConfig) -> None:
    # Insert after the Title paragraph, before the first chapter heading.
    subtitle = document.add_paragraph(config.subtitle, style="WSC Subtitle")
    _set_paragraph_border(subtitle, side="bottom", color=GOLD, size=10, space=8)


def _reorder_subtitle(document: DocumentType) -> None:
    title = next((p for p in document.paragraphs if p.style.name == "Title"), None)
    subtitle = next(
        (p for p in document.paragraphs if p.style.name == "WSC Subtitle"), None
    )
    if title is None or subtitle is None:
        return
    title._p.addnext(subtitle._p)


def _set_core_properties(document: DocumentType, title: str) -> None:
    properties = document.core_properties
    properties.title = title
    properties.subject = "World Scholar's Cup 2026 long-form curriculum guide"
    properties.author = "WSCapp"
    properties.last_modified_by = "WSCapp"
    properties.keywords = "WSCapp, World Scholar's Cup, 2026, study guide"
    properties.comments = "Generated from the semantic guide.html source."
    properties.created = FIXED_TIMESTAMP
    properties.modified = FIXED_TIMESTAMP


def _normalize_docx_archive(path: Path) -> None:
    """Give DOCX members stable timestamps and ordering for cleaner diffs."""

    temporary = path.with_suffix(".normalized.docx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
        temporary, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, ZIP_TIMESTAMP)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, source.read(name))
    temporary.replace(path)


def build_docx(config: GuideConfig, output_path: Path) -> Path:
    source_path = THEME_SECTIONS / config.slug / "guide.html"
    if not source_path.exists():
        raise FileNotFoundError(source_path)
    if not GUIDE_MASCOT.exists():
        raise FileNotFoundError(GUIDE_MASCOT)

    root = html.fragment_fromstring(
        source_path.read_text(encoding="utf-8"), create_parent="div"
    )
    document = Document()
    _remove_initial_empty_paragraph(document)
    _configure_styles(document)
    _add_page_background(document)
    _configure_page(document, config)

    h3_count = [0]
    for child in root:
        _walk_block(document, child, h3_count=h3_count)
    if not any(p.style.name == "WSC Subtitle" for p in document.paragraphs):
        _add_subtitle(document, config)
    _reorder_subtitle(document)

    title_element = root.xpath(".//h2[1]")
    title = title_element[0].text_content().strip() if title_element else config.footer_title
    _set_core_properties(document, title)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    _normalize_docx_archive(output_path)
    return output_path


def _find_soffice() -> Path:
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if executable:
        return Path(executable)
    if BUNDLED_SOFFICE.exists():
        return BUNDLED_SOFFICE
    raise FileNotFoundError(
        "LibreOffice was not found. Install it or make `soffice` available on PATH."
    )


def convert_to_pdf(docx_path: Path, output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    expected = output_dir / f"{docx_path.stem}.pdf"
    if expected.exists():
        expected.unlink()

    soffice = _find_soffice()
    with tempfile.TemporaryDirectory(prefix="wscapp-lo-") as profile_dir:
        profile_uri = Path(profile_dir).resolve().as_uri()
        filter_options = (
            'pdf:writer_pdf_Export:{"UseTaggedPDF":{"type":"boolean","value":"true"},'
            '"ExportBookmarks":{"type":"boolean","value":"true"},'
            '"ExportFormFields":{"type":"boolean","value":"true"}}'
        )
        environment = os.environ.copy()
        environment.setdefault("SOURCE_DATE_EPOCH", "1783987200")
        result = subprocess.run(
            [
                str(soffice),
                f"-env:UserInstallation={profile_uri}",
                "--headless",
                "--convert-to",
                filter_options,
                "--outdir",
                str(output_dir.resolve()),
                str(docx_path.resolve()),
            ],
            check=False,
            capture_output=True,
            text=True,
            env=environment,
        )
    if result.returncode != 0 or not expected.exists():
        details = "\n".join(part for part in (result.stdout, result.stderr) if part)
        raise RuntimeError(f"LibreOffice PDF export failed for {docx_path}:\n{details}")
    return expected


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build long-form WSCapp DOCX and PDF guides from guide.html."
    )
    parser.add_argument(
        "--guide",
        choices=("all", *GUIDE_ALIASES, *GUIDES),
        default="all",
        help="Build all guides, a full section slug, or the call/draft alias.",
    )
    parser.add_argument(
        "--skip-pdf",
        action="store_true",
        help="Build only the editable DOCX delivery files.",
    )
    parser.add_argument(
        "--output-root",
        type=Path,
        default=DEFAULT_OUTPUT_ROOT,
        help=(
            "Root containing docx/ and pdf/ output folders "
            f"(default: {DEFAULT_OUTPUT_ROOT})."
        ),
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if args.guide == "all":
        selected = GUIDES
    else:
        guide_key = GUIDE_ALIASES.get(args.guide, args.guide)
        selected = {guide_key: GUIDES[guide_key]}
    output_root = args.output_root.resolve()

    for key, config in selected.items():
        docx_path = build_docx(config, output_root / "docx" / f"{config.slug}.docx")
        print(f"{key}: DOCX -> {docx_path}")
        if not args.skip_pdf:
            pdf_path = convert_to_pdf(docx_path, output_root / "pdf")
            print(f"{key}: PDF  -> {pdf_path}")


if __name__ == "__main__":
    main()
